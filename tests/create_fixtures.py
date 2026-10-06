#!/usr/bin/env python3
import os
import sys

# Reconfigure stdout for UTF-8 on Windows
if sys.platform == 'win32':
    try:
        sys.stdout.reconfigure(encoding='utf-8', errors='replace')
    except Exception:
        pass

os.makedirs('tests/fixtures', exist_ok=True)

# 1. Create sample_resume.docx
import docx
doc = docx.Document()
doc.add_heading('Jane Doe', 0)
doc.add_paragraph('jane.doe.systems@gmail.com | +1 (555) 432-8765 | San Francisco, CA')
doc.add_paragraph('https://linkedin.com/in/janedoe-systems | https://github.com/janedoe-infra')
doc.add_heading('Summary', level=1)
doc.add_paragraph('Principal Systems & Infrastructure Engineer with 8+ years experience designing, deploying, and maintaining high-scale distributed systems and core cloud infrastructure.')
doc.add_heading('Education', level=1)
doc.add_paragraph('Master of Science in Computer Science, Stanford University (2018)')
doc.add_heading('Experience', level=1)
doc.add_paragraph('2018 - Present: Lead Systems Engineer at CloudTech. Architected distributed systems in Rust, Python, Go.')
doc.add_heading('Skills', level=1)
doc.add_paragraph('Rust, Python, Go, C++, Distributed Systems, Systems Engineering, Core Infrastructure, Kubernetes, Docker, Linux, AWS, GCP, Terraform, CI/CD, gRPC')
doc.add_heading('Work Authorization', level=1)
doc.add_paragraph('US Citizen (No sponsorship required)')
doc.save('tests/fixtures/sample_resume.docx')
print('Created tests/fixtures/sample_resume.docx')

# 2. Create sample_resume.pdf
text_lines = [
    "Jane Doe",
    "jane.doe.systems@gmail.com | +1 (555) 432-8765 | San Francisco, CA",
    "https://linkedin.com/in/janedoe-systems | https://github.com/janedoe-infra",
    "Summary: Principal Systems and Infrastructure Engineer with 8+ years experience in distributed systems.",
    "Education: Master of Science in Computer Science, Stanford University",
    "Experience: 2018 - Present: Lead Systems Engineer at CloudTech in Rust, Python, Go.",
    "Skills: Rust, Python, Go, Kubernetes, Docker, Distributed Systems, Systems Engineering, Linux, AWS, gRPC",
    "Work Authorization: US Citizen (No sponsorship required)"
]

stream_content = "BT\n/F1 12 Tf\n72 720 Td\n18 TL\n"
for line in text_lines:
    safe = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream_content += f"({safe}) '\n"
stream_content += "ET\n"

stream_bytes = stream_content.encode("latin-1")
stream_len = len(stream_bytes)

# Build simple PDF
pdf_parts = []
pdf_parts.append(b"%PDF-1.4\n")

offsets = [0]
def add_obj(data):
    offsets.append(sum(len(p) for p in pdf_parts))
    pdf_parts.append(data)

add_obj(b"1 0 obj\n<< /Type /Catalog /Pages 2 0 R >>\nendobj\n")
add_obj(b"2 0 obj\n<< /Type /Pages /Kids [3 0 R] /Count 1 >>\nendobj\n")
add_obj(b"3 0 obj\n<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>\nendobj\n")
add_obj(f"4 0 obj\n<< /Length {stream_len} >>\nstream\n".encode("ascii") + stream_bytes + b"\nendstream\nendobj\n")
add_obj(b"5 0 obj\n<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>\nendobj\n")

xref_pos = sum(len(p) for p in pdf_parts)
pdf_parts.append(f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode("ascii"))
for off in offsets[1:]:
    pdf_parts.append(f"{off:010d} 00000 n \n".encode("ascii"))

pdf_parts.append(f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref_pos}\n%%EOF\n".encode("ascii"))

with open("tests/fixtures/sample_resume.pdf", "wb") as f:
    f.write(b"".join(pdf_parts))
print("Created tests/fixtures/sample_resume.pdf")
